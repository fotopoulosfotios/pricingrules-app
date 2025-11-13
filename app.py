import io, os, re, json, yaml
import pandas as pd
import streamlit as st
import chardet
from pdfminer.high_level import extract_text as pdf_extract_text
from docx import Document

# ---------------------------------------------------------
#  Ρυθμίσεις εφαρμογής
# ---------------------------------------------------------
st.set_page_config(page_title="Έλεγχος Τιμολογήσεων", page_icon="💶", layout="wide")
st.title("💶 Σύστημα Ελέγχου Τιμολογήσεων Εισερχόμενων Εντολών")

# ---------------------------------------------------------
#  Απαιτούμενες στήλες
# ---------------------------------------------------------
REQUIRED_COLS = ["BIC", "Currency", "Amount", "ChargedAmount"]

# Συνώνυμα για αυτόματο mapping
SYNONYMS = {
    "bic": ["bic", "sender_bic", "bank_bic"],
    "currency": ["currency", "ccy", "curr"],
    "amount": ["amount", "amt", "value"],
    "chargedamount": ["chargedamount", "charged_amount", "fee", "charges", "pricing"],
}

# ---------------------------------------------------------
#  Validate κανόνων YAML
# ---------------------------------------------------------
def validate_rules(rules):
    required_fields = {"id", "description", "condition", "action"}

    for r in rules:
        if not isinstance(r, dict):
            return "Κάποιος κανόνας δεν είναι αντικείμενο."

        missing = required_fields - r.keys()
        if missing:
            return f"Λείπουν πεδία: {', '.join(missing)}"

        if "field" not in r["condition"]:
            return "Κάποιος κανόνας δεν έχει condition.field"

        if "set_charge" not in r["action"]:
            return "Κάποιος κανόνας δεν έχει action.set_charge"

    return None

# ---------------------------------------------------------
#  Κανονικοποίηση στηλών
# ---------------------------------------------------------
def normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    col_map = {}
    lower_cols = {c.lower(): c for c in df.columns}

    for target, alts in SYNONYMS.items():
        for a in alts:
            if a in lower_cols:
                col_map[lower_cols[a]] = target.capitalize() if target != "chargedamount" else "ChargedAmount"
                break

    for c in df.columns:
        if c in REQUIRED_COLS:
            col_map[c] = c

    return df.rename(columns=col_map)

def find_missing_columns(df: pd.DataFrame):
    return [c for c in REQUIRED_COLS if c not in df.columns]

# ---------------------------------------------------------
#  Ανάγνωση αρχείων
# ---------------------------------------------------------
def read_csv_file(file) -> pd.DataFrame:
    raw = file.read()
    enc = chardet.detect(raw).get("encoding") or "utf-8"
    return pd.read_csv(io.BytesIO(raw), encoding=enc)

def read_xlsx_file(file) -> pd.DataFrame:
    return pd.read_excel(file)

def read_json_file(file) -> pd.DataFrame:
    try:
        obj = json.load(file)
    except:
        obj = json.loads(file.read().decode("utf-8", errors="ignore"))
    if isinstance(obj, list):
        return pd.DataFrame(obj)
    if isinstance(obj, dict):
        for v in obj.values():
            if isinstance(v, list) and all(isinstance(x, dict) for x in v):
                return pd.DataFrame(v)
        return pd.DataFrame([obj])
    return pd.DataFrame()

def read_txt_file(file) -> pd.DataFrame:
    raw = file.read()
    enc = chardet.detect(raw).get("encoding") or "utf-8"
    text = raw.decode(enc, errors="ignore")
    rows = [re.split(r"[;,|\t]| {2,}", ln.strip()) for ln in text.splitlines() if ln.strip()]
    return pd.DataFrame(rows)

def read_pdf_file(file) -> pd.DataFrame:
    try:
        text = pdf_extract_text(file)
        rows = [re.split(r"[;,|\t]| {2,}", ln.strip()) for ln in text.splitlines() if ln.strip()]
        return pd.DataFrame(rows)
    except Exception:
        st.error("❌ Αποτυχία ανάγνωσης PDF")
        return pd.DataFrame()

def read_docx_file(file) -> pd.DataFrame:
    byts = file.read()
    doc = Document(io.BytesIO(byts))
    tables = doc.tables
    if not tables:
        return pd.DataFrame()
    tbl = tables[0]
    rows = [[c.text.strip() for c in r.cells] for r in tbl.rows]
    header = rows[0]
    body = rows[1:]
    df = pd.DataFrame(body, columns=header)
    return df

# ---------------------------------------------------------
#  OCR (προαιρετικό)
# ---------------------------------------------------------
def read_scanned_pdf_with_ocr(file_bytes, filename):
    api_key = st.secrets.get("OCRSPACE_API_KEY", None)
    if not api_key:
        st.warning("🔍 Το PDF ίσως είναι σκαναρισμένο. Προσθέστε OCR API key στα secrets.")
        return pd.DataFrame()

    import requests
    files = {"file": (filename, io.BytesIO(file_bytes))}
    data = {"apikey": api_key, "language": "eng"}

    try:
        r = requests.post("https://api.ocr.space/parse/image", files=files, data=data)
        js = r.json()
        text = "\n".join([p["ParsedText"] for p in js.get("ParsedResults", [])])
        rows = [re.split(r"[;,|\t]| {2,}", ln.strip()) for ln in text.splitlines() if ln.strip()]
        return pd.DataFrame(rows)
    except Exception as e:
        st.error(f"❌ OCR Error: {e}")
        return pd.DataFrame()

# ---------------------------------------------------------
#  Φόρτωμα αρχείου εντολών
# ---------------------------------------------------------
def load_uploaded_file(file):
    name = file.name.lower()

    if name.endswith(".csv"):
        return read_csv_file(file)
    if name.endswith(".xlsx"):
        return read_xlsx_file(file)
    if name.endswith(".json"):
        return read_json_file(file)
    if name.endswith(".txt"):
        return read_txt_file(file)
    if name.endswith(".pdf"):
        df = read_pdf_file(file)
        if df.empty or len(df.columns) <= 2:
            file.seek(0)
            df = read_scanned_pdf_with_ocr(file.read(), file.name)
        return df
    if name.endswith(".docx"):
        return read_docx_file(file)

    st.error("❌ Μη υποστηριζόμενος τύπος αρχείου.")
    return pd.DataFrame()

# ---------------------------------------------------------
#  Εφαρμογή κανόνων
# ---------------------------------------------------------
def apply_rules(df: pd.DataFrame, rules):
    output = []

    for i, row in df.iterrows():
        rule_applied = None
        charge = 0.0

        for rule in rules:
            cond = rule["condition"]
            field = cond.get("field")
            if field != "Amount":
                continue

            try:
                amount = float(row.get("Amount", 0))
            except:
                amount = 0

            matched = False

            if "less_equal" in cond and amount <= cond["less_equal"]:
                matched = True
            elif "between" in cond:
                lo, hi = cond["between"]
                if lo <= amount <= hi:
                    matched = True
            elif "greater_than" in cond and amount > cond["greater_than"]:
                matched = True

            if matched:
                charge = rule["action"]["set_charge"]
                rule_applied = rule["id"]
                break

        try:
            charged_actual = float(row.get("ChargedAmount", 0))
        except:
            charged_actual = 0.0

        diff = charged_actual - charge

        output.append({
            "BIC": row.get("BIC", ""),
            "Currency": row.get("Currency", ""),
            "Amount": row.get("Amount", ""),
            "ChargedAmount": charged_actual,
            "ExpectedCharge": charge,
            "Difference": diff,
            "RuleApplied": rule_applied or "—"
        })

    return pd.DataFrame(output)

# ---------------------------------------------------------
#  Sidebar — Αρχείο κανόνων
# ---------------------------------------------------------
st.sidebar.header("⚙️ Αρχείο Κανόνων")
rules_file = st.sidebar.file_uploader("Ανέβασε YAML κανόνων", type=["yaml", "yml"])

if rules_file:
    try:
        y = yaml.safe_load(rules_file)
        rules = y.get("rules", [])
        err = validate_rules(rules)
        if err:
            st.sidebar.error(f"❌ Σφάλμα στους κανόνες: {err}")
            rules = []
        else:
            st.sidebar.success("✅ Κανόνες φορτώθηκαν.")
    except Exception as e:
        st.sidebar.error(f"❌ Μη έγκυρο YAML: {e}")
        rules = []
else:
    rules = []
    st.sidebar.info("🔹 Ανέβασε YAML κανόνων για να γίνει έλεγχος.")

# ---------------------------------------------------------
#  Upload αρχείου τιμολογημένων εντολών
# ---------------------------------------------------------
st.header("📂 Ανέβασε Αρχείο Τιμολογημένων Εντολών")

uploaded_file = st.file_uploader(
    "Υποστηριζόμενα: CSV, XLSX, PDF, DOCX, TXT, JSON",
    type=["csv", "xlsx", "pdf", "docx", "txt", "json"]
)

df = pd.DataFrame()

if uploaded_file:
    df = load_uploaded_file(uploaded_file)

    if df.empty:
        st.error("❌ Δεν βρέθηκαν δεδομένα.")
    else:
        st.success(f"✅ Φορτώθηκαν {len(df)} γραμμές.")
        st.subheader("📄 Προεπισκόπηση")
        st.dataframe(df.head(50), use_container_width=True)

        df = normalize_columns(df)

        missing = find_missing_columns(df)
        if missing:
            st.error(f"❌ Λείπουν στήλες: {', '.join(missing)}")
        else:
            if not rules:
                st.warning("⚠️ Δεν υπάρχουν κανόνες.")
            else:
                st.header("📊 Αποτελέσματα Ελέγχου")
                result = apply_rules(df, rules)
                st.dataframe(result, use_container_width=True)

                st.subheader("📈 Στατιστικά")
                total_diff = result["Difference"].sum()
                ok = (result["Difference"] == 0).sum()
                wrong = (result["Difference"] != 0).sum()

                st.metric("✔️ Σωστά", ok)
                st.metric("❌ Λάθος", wrong)
                st.metric("Σύνολο Απόκλισης", total_diff)

                st.table(
                    result["RuleApplied"].value_counts().reset_index().rename(
                        columns={"index": "Rule", "RuleApplied": "Count"}
                    )
                )

                # ---------------------------------------------------------
                #  Exports
                # ---------------------------------------------------------
                st.subheader("📥 Λήψη Αποτελεσμάτων")

                # Excel
                out_xlsx = io.BytesIO()
                result.to_excel(out_xlsx, index=False)
                st.download_button(
                    "📥 Λήψη Excel",
                    out_xlsx.getvalue(),
                    "audit_results.xlsx",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )

                # PDF
                try:
                    from fpdf import FPDF
                    pdf = FPDF()
                    pdf.add_page()
                    pdf.set_font("Arial", size=10)

                    for _, row in result.iterrows():
                        line = f"{row['BIC']} | {row['Amount']} | Charged={row['ChargedAmount']} | Expected={row['ExpectedCharge']} | Diff={row['Difference']}"
                        pdf.cell(0, 5, txt=line, ln=1)

                    out_pdf = io.BytesIO(pdf.output(dest="S").encode("latin1"))

                    st.download_button(
                        "📥 Λήψη PDF",
                        out_pdf.getvalue(),
                        "audit_results.pdf",
                        "application/pdf"
                    )
                except:
                    st.info("Προσθέστε στο requirements: fpdf==1.7.2")

st.markdown("---")
st.caption("© 2025 Σύστημα Ελέγχου Τιμολογήσεων — Υποστήριξη αρχείων: CSV, XLSX, PDF(text), DOCX, TXT, JSON.")
